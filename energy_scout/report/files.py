"""Экспорт в CSV, Excel и автономную HTML-страницу с фильтрами."""
from __future__ import annotations

import csv
import html
import json
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..models import Product
from ..pipeline import RunReport

HEADERS = [
    "Редкость", "Метки", "Название", "Бренд", "Цена, тг", "Старая цена, тг",
    "Скидка, %", "Объём, л", "Цена за литр, тг", "Наличие", "Магазин",
    "Город", "Источник", "Заметка", "Ссылка",
]


def _row_values(p: Product) -> list:
    return [
        p.rarity_tier,
        ", ".join(p.tags),
        p.name,
        p.brand or "",
        p.price or "",
        p.old_price or "",
        p.discount_percent or "",
        p.volume_l or "",
        p.price_per_l or "",
        {True: "есть", False: "нет", None: "н/д"}[p.in_stock],
        p.shop_name or "",
        p.city or "",
        p.source,
        p.note or "",
        p.url or "",
    ]


def write_csv(path: Path, products: list[Product]) -> None:
    with path.open("w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.writer(fh, delimiter=";")
        writer.writerow(HEADERS)
        for p in products:
            writer.writerow(_row_values(p))


def write_xlsx(path: Path, report: RunReport) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Результаты"
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    tier_fill = {
        "exclusive": PatternFill("solid", fgColor="F4D7FF"),
        "rare": PatternFill("solid", fgColor="FFF3C4"),
    }
    for p in report.products:
        ws.append(_row_values(p))
        if p.rarity_tier in tier_fill:
            for cell in ws[ws.max_row]:
                cell.fill = tier_fill[p.rarity_tier]

    widths = [11, 26, 42, 14, 10, 12, 9, 9, 13, 9, 26, 10, 11, 32, 40]
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    summary = wb.create_sheet("Сводка")
    stats = report.stats()
    summary.append(["Параметр", "Значение"])
    summary.append(["Сформировано", datetime.now().strftime("%Y-%m-%d %H:%M")])
    summary.append(["Всего товаров", stats["total"]])
    summary.append(["Эксклюзивных", stats["exclusive"]])
    summary.append(["Редких", stats["rare"]])
    summary.append(["С ценой", stats["with_price"]])
    summary.append(["В наличии", stats["in_stock"]])
    summary.append([])
    summary.append(["Источник", "Товаров в выдаче", "Ошибки"])
    for r in report.source_results:
        summary.append([r.label, len(r.products), "; ".join(r.errors)])
    for cell in summary[1]:
        cell.font = Font(bold=True)

    wb.save(path)


HTML_TEMPLATE = """<!doctype html>
<html lang="ru"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Energy Scout — __CITY__</title>
<style>
:root { --bg:#0f1220; --card:#1a1f35; --txt:#e8ebf7; --mut:#9aa3c0;
--acc:#6c8cff; --rare:#ffd166; --excl:#e0aaff; --ok:#4ade80; --no:#f87171; }
* { box-sizing:border-box; }
body { margin:0; font:15px/1.45 system-ui,Segoe UI,Roboto,sans-serif;
background:var(--bg); color:var(--txt); padding:24px; }
h1 { margin:0 0 4px; font-size:24px; }
.sub { color:var(--mut); margin-bottom:18px; }
.filters { display:flex; flex-wrap:wrap; gap:10px; margin-bottom:18px; }
input,select,button { background:var(--card); color:var(--txt);
border:1px solid #2c3354; border-radius:10px; padding:9px 12px; font-size:14px; }
input[type=search] { min-width:260px; flex:1; }
button { cursor:pointer; }
.card { background:var(--card); border:1px solid #252b45; border-radius:14px;
padding:14px 16px; margin-bottom:12px; display:grid;
grid-template-columns:1fr auto; gap:6px 16px; }
.card h3 { margin:0; font-size:16px; grid-column:1/-1; }
.price { font-size:20px; font-weight:700; white-space:nowrap; }
.old { color:var(--mut); text-decoration:line-through; font-size:13px; font-weight:400; }
.meta { color:var(--mut); font-size:13px; }
.tag { display:inline-block; padding:2px 8px; border-radius:999px; font-size:12px;
margin:2px 4px 2px 0; background:#262d4d; }
.tag.exclusive { background:#3b2354; color:var(--excl); }
.tag.rare { background:#4a3d17; color:var(--rare); }
.tag.ok { color:var(--ok); } .tag.no { color:var(--no); }
a { color:var(--acc); text-decoration:none; }
a:hover { text-decoration:underline; }
#empty { color:var(--mut); padding:30px 0; }
</style></head><body>
<h1>⚡ Energy Scout — __CITY__</h1>
<div class="sub">Сформировано __GENERATED__. Источники: Dina, Magnum, Anvar, веб-поиск.</div>
<div class="filters">
  <input id="q" type="search" placeholder="Поиск по названию, бренду, метке…">
  <select id="rarity">
    <option value="">Любая редкость</option>
    <option value="exclusive">Эксклюзив</option>
    <option value="rare">Редкое</option>
    <option value="ordinary">Обычные</option>
  </select>
  <select id="stock">
    <option value="">Любое наличие</option>
    <option value="in">В наличии</option>
    <option value="out">Нет в наличии</option>
  </select>
  <input id="maxprice" type="number" placeholder="Макс. цена, тг" min="0">
  <button id="reset">Сбросить</button>
</div>
<div id="count" class="meta"></div>
<div id="list"></div>
<div id="empty" hidden>Ничего не подходит под фильтры.</div>
<script>
const DATA = __DATA__;
const list = document.getElementById('list'), empty = document.getElementById('empty'),
      countEl = document.getElementById('count');
const tierName = {exclusive:'ЭКСКЛЮЗИВ', rare:'РЕДКОЕ', ordinary:'обычный'};
function card(p) {
  const price = p.price ? `${p.price.toLocaleString('ru-RU')} тг` : 'цена не указана';
  const old = p.old_price ? `<span class="old">${p.old_price.toLocaleString('ru-RU')} тг</span>` : '';
  const perL = p.price_per_l ? `<span class="meta">${p.price_per_l.toLocaleString('ru-RU')} тг/л</span>` : '';
  const stock = p.in_stock === true ? '<span class="tag ok">в наличии</span>'
    : p.in_stock === false ? '<span class="tag no">нет</span>'
    : '<span class="tag">наличие н/д</span>';
  const tags = (p.tags||[]).map(t => `<span class="tag">${t}</span>`).join('');
  const link = p.url ? ` <a href="${p.url}" target="_blank" rel="noopener">открыть ↗</a>` : '';
  return `<div class="card" data-search="${(p.name+' '+(p.brand||'')+' '+(p.tags||[]).join(' ')).toLowerCase()}"
      data-tier="${p.rarity_tier}" data-stock="${p.in_stock}" data-price="${p.price ?? ''}">
    <h3><span class="tag ${p.rarity_tier}">${tierName[p.rarity_tier]}</span> ${p.name}</h3>
    <div><div class="price">${price} ${old}</div>${perL}</div>
    <div style="text-align:right">${stock}<div class="meta">${p.shop_name||''}</div></div>
    <div class="meta" style="grid-column:1/-1">${tags} ${p.note||''}${link}</div>
  </div>`;
}
function render() {
  const q = document.getElementById('q').value.trim().toLowerCase();
  const rarity = document.getElementById('rarity').value;
  const stock = document.getElementById('stock').value;
  const max = Number(document.getElementById('maxprice').value) || Infinity;
  let shown = 0;
  for (const el of list.children) {
    const okQ = !q || el.dataset.search.includes(q);
    const okR = !rarity || el.dataset.tier === rarity;
    const okS = !stock || (stock === 'in' && el.dataset.stock === 'true')
      || (stock === 'out' && el.dataset.stock === 'false');
    const okP = el.dataset.price === '' || Number(el.dataset.price) <= max;
    const visible = okQ && okR && okS && okP;
    el.hidden = !visible;
    if (visible) shown++;
  }
  empty.hidden = shown !== 0;
  countEl.textContent = `Показано ${shown} из ${DATA.length}`;
}
list.innerHTML = DATA.map(card).join('');
['q','rarity','stock','maxprice'].forEach(id =>
  document.getElementById(id).addEventListener('input', render));
document.getElementById('reset').onclick = () => {
  ['q','maxprice'].forEach(id => document.getElementById(id).value = '');
  document.getElementById('rarity').value = '';
  document.getElementById('stock').value = '';
  render();
};
render();
</script></body></html>"""


def write_html(path: Path, products: list[Product], city: str) -> None:
    payload = json.dumps([p.to_dict() for p in products], ensure_ascii=False)
    text = (
        HTML_TEMPLATE
        .replace("__CITY__", html.escape(city))
        .replace("__GENERATED__", datetime.now().strftime("%Y-%m-%d %H:%M"))
        .replace("__DATA__", payload)
    )
    path.write_text(text, encoding="utf-8")
